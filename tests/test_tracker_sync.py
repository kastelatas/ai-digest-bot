import importlib.util
import sqlite3
import tempfile
import unittest
from contextlib import closing
from datetime import datetime, timezone
from pathlib import Path

from digest_bot.ads import book_ad
from digest_bot.db import Database
from digest_bot.models import AdStatus, Draft, DraftStatus
from tests.helpers import make_item

try:
    from openpyxl import load_workbook

    from digest_bot.tracker_sync import TrackerSyncError, sync_tracker

    HAVE_OPENPYXL = True
except ImportError:  # sync-tracker — единственная часть, которой нужен openpyxl
    HAVE_OPENPYXL = False

TRACKER_DIR = Path(__file__).resolve().parent.parent / "tracker"


def _build_template(path: Path) -> None:
    """Свежий шаблон трекера тем же скриптом, которым сделан настоящий, — чтобы тест
    не зависел от данных, которые пользователь уже вписал в tracker/channel_tracker.xlsx."""
    spec = importlib.util.spec_from_file_location("build_tracker", TRACKER_DIR / "build_tracker.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    module.build().save(path)


def utc(y, m, d, hh=12, mm=0):
    return datetime(y, m, d, hh, mm, tzinfo=timezone.utc)


@unittest.skipUnless(HAVE_OPENPYXL, "нужен openpyxl (pip install -r requirements.txt)")
class SyncTrackerTests(unittest.TestCase):
    def setUp(self):
        self._tmpdir = tempfile.TemporaryDirectory()
        tmp = Path(self._tmpdir.name)
        self.db = Database(tmp / "test.db")
        self.xlsx = tmp / "tracker.xlsx"
        _build_template(self.xlsx)

    def tearDown(self):
        self._tmpdir.cleanup()

    # ---------- хелперы ----------

    def _stat(self, subscribers, captured_at: datetime):
        """Замер подписчиков с заданным временем (в БД оно хранится как naive UTC)."""
        self.db.record_channel_stat(subscribers)
        with closing(sqlite3.connect(self.db.path)) as conn, conn:
            conn.execute(
                "UPDATE channel_stats SET captured_at = ? WHERE id = (SELECT MAX(id) FROM channel_stats)",
                (captured_at.replace(tzinfo=None).isoformat(),),
            )

    def _publish(self, guid, published_at, message_id, views=None):
        self.db.add_item(make_item(guid=guid, title=guid, link=f"https://example.com/{guid}"))
        draft_id = self.db.create_draft(
            Draft(id=None, item_guid=guid, source_name="Src", title=guid, link=f"https://example.com/{guid}",
                  draft_text="text", tags=[], status=DraftStatus.PENDING, created_at=published_at)
        )
        self.db.mark_published(draft_id, message_id, published_at)
        if views is not None:
            self.db.record_metric_snapshot(message_id, views=views, forwards=0, reactions=0)

    def _sync(self, tz="UTC"):
        return sync_tracker(self.db, self.xlsx, tz)

    def _wb(self):
        return load_workbook(self.xlsx)

    @staticmethod
    def _row(ws, r, cols):
        return [ws.cell(row=r, column=c).value for c in cols]

    # ---------- Метрики ----------

    def test_metrics_row_appended_after_examples_with_formulas_intact(self):
        before = self._wb()["Метрики"]
        f_c7, f_f7 = before["C7"].value, before["F7"].value
        example_row_6 = self._row(before, 6, range(1, 8))

        self._stat(1100, utc(2026, 9, 22, 8))
        self._publish("g1", utc(2026, 9, 22, 9), 101, views=400)
        self._publish("g2", utc(2026, 9, 22, 10), 102, views=200)
        result = self._sync()

        ws = self._wb()["Метрики"]
        self.assertEqual(result.metrics_added, 1)
        # 5–6 — строки-примеры из шаблона, первая пустая строка-заготовка — 7-я
        self.assertEqual(self._row(ws, 7, (1, 2, 4, 5)), [datetime(2026, 9, 22), 1100, 2, 300])
        self.assertEqual(ws["C7"].value, f_c7)  # формулы не тронуты
        self.assertEqual(ws["F7"].value, f_f7)
        self.assertEqual(self._row(ws, 6, range(1, 8)), example_row_6)
        self.assertIsNone(ws["A8"].value)

    def test_example_row_dated_same_day_does_not_block_that_day(self):
        # шаблон строится «сегодня» и кладёт пример (голубая заливка) с сегодняшней датой
        example_date = self._wb()["Метрики"]["A6"].value
        self.assertIsInstance(example_date, datetime)
        self._stat(2, example_date.replace(hour=12, tzinfo=timezone.utc))

        result = self._sync()

        ws = self._wb()["Метрики"]
        self.assertEqual(result.metrics_added, 1)
        self.assertEqual(self._row(ws, 7, (1, 2)), [example_date, 2])
        self.assertEqual(ws["B6"].value, 1012)  # сам пример не тронут
        self.assertEqual(self._sync().metrics_added, 0)  # и повторно день не дублируется

    def test_second_run_does_not_duplicate_and_new_day_goes_to_next_row(self):
        self._stat(1100, utc(2026, 9, 22, 8))
        self._sync()
        self.assertEqual(self._sync().metrics_added, 0)

        self._stat(1120, utc(2026, 9, 23, 8))
        self.assertEqual(self._sync().metrics_added, 1)

        ws = self._wb()["Метрики"]
        self.assertEqual(self._row(ws, 7, (1, 2)), [datetime(2026, 9, 22), 1100])
        self.assertEqual(self._row(ws, 8, (1, 2)), [datetime(2026, 9, 23), 1120])
        self.assertIsNone(ws["A9"].value)

    def test_all_missing_days_synced_in_one_run_in_date_order_last_snapshot_wins(self):
        self._stat(1300, utc(2026, 9, 24, 8))  # вставлены не по порядку
        self._stat(1100, utc(2026, 9, 22, 8))
        self._stat(1110, utc(2026, 9, 22, 20))  # второй замер того же дня — он и нужен
        self._stat(1200, utc(2026, 9, 23, 8))

        result = self._sync()

        ws = self._wb()["Метрики"]
        self.assertEqual(result.metrics_added, 3)
        self.assertEqual(
            [self._row(ws, r, (1, 2)) for r in (7, 8, 9)],
            [[datetime(2026, 9, 22), 1110], [datetime(2026, 9, 23), 1200], [datetime(2026, 9, 24), 1300]],
        )

    def test_day_already_in_xlsx_is_skipped_not_overwritten(self):
        wb = self._wb()
        wb["Метрики"]["A7"] = datetime(2026, 9, 22)
        wb["Метрики"]["B7"] = 999
        wb.save(self.xlsx)

        self._stat(1100, utc(2026, 9, 22, 8))
        self.assertEqual(self._sync().metrics_added, 0)
        self.assertEqual(self._wb()["Метрики"]["B7"].value, 999)

    def test_posts_counted_per_day_and_reach_blank_without_views(self):
        self._stat(1000, utc(2026, 9, 22, 8))
        self._stat(1010, utc(2026, 9, 23, 8))
        self._publish("g1", utc(2026, 9, 22, 9), 101)  # без просмотров (нет telethon)
        self._publish("g2", utc(2026, 9, 22, 15), 102)
        self._publish("g3", utc(2026, 9, 23, 9), 103)

        self._sync()

        ws = self._wb()["Метрики"]
        self.assertEqual(self._row(ws, 7, (4, 5)), [2, None])
        self.assertEqual(self._row(ws, 8, (4, 5)), [1, None])

    def test_day_boundary_uses_channel_timezone(self):
        # 22:30 UTC 22 сентября — это уже 23 сентября 01:30 в Киеве (UTC+3)
        self._stat(1100, utc(2026, 9, 22, 22, 30))
        self._sync("Europe/Kyiv")
        self.assertEqual(self._wb()["Метрики"]["A7"].value, datetime(2026, 9, 23))

    def test_no_free_template_rows_reports_skipped_and_keeps_file_valid(self):
        wb = self._wb()
        ws = wb["Метрики"]
        for r in range(5, ws.max_row + 1):
            if isinstance(ws.cell(row=r, column=6).value, str) and ws.cell(row=r, column=6).value.startswith("="):
                ws.cell(row=r, column=1, value=datetime(2020, 1, 1))
                ws.cell(row=r, column=2, value=1)
        wb.save(self.xlsx)

        self._stat(1100, utc(2026, 9, 22, 8))
        result = self._sync()

        self.assertEqual(result.metrics_added, 0)
        self.assertEqual(result.skipped_no_room, 1)
        self.assertIsNotNone(self._wb()["Дашборд"]["B4"].value)  # файл цел

    # ---------- Продажи_рекламы ----------

    def test_ad_appended_after_example_with_id_and_formula_intact(self):
        f_h6 = self._wb()["Продажи_рекламы"]["H6"].value
        ad_id = book_ad(self.db, advertiser="Beta Corp", price=250.5, scheduled_at=utc(2026, 9, 25, 10),
                        duration_hours=24, currency="EUR", contact="@beta", erid="ERID-1", notes="счёт выставлен")

        result = self._sync()

        ws = self._wb()["Продажи_рекламы"]
        self.assertEqual(result.ads_added, 1)
        # строка 5 — пример из шаблона, первая пустая строка-заготовка — 6-я
        self.assertEqual(
            self._row(ws, 6, (1, 2, 3, 4, 5, 6, 9, 10, 11, 12)),
            [datetime(2026, 9, 25), "Beta Corp", "@beta", "Пост 24ч", 250.5, "EUR",
             "ERID-1", "Забронировано", "счёт выставлен", ad_id],
        )
        self.assertEqual(ws["L4"].value, "ID брони")
        self.assertEqual(ws["H6"].value, f_h6)

    def test_ads_second_run_does_not_duplicate_and_only_new_ad_is_added(self):
        book_ad(self.db, advertiser="A", price=10, scheduled_at=utc(2026, 9, 25), duration_hours=48)
        self._sync()
        self.assertEqual(self._sync().ads_added, 0)

        second = book_ad(self.db, advertiser="B", price=20, scheduled_at=utc(2026, 9, 26), duration_hours=48)
        self.assertEqual(self._sync().ads_added, 1)

        ws = self._wb()["Продажи_рекламы"]
        self.assertEqual([ws.cell(row=r, column=2).value for r in (6, 7)], ["A", "B"])
        self.assertEqual(ws["L7"].value, second)
        self.assertIsNone(ws["B8"].value)

    def test_ad_status_and_views_mapped(self):
        ad_id = book_ad(self.db, advertiser="A", price=10, scheduled_at=utc(2026, 9, 25), duration_hours=48)
        self.db.set_ad_channel_message(ad_id, 777)
        self.db.mark_ad_removed(ad_id)
        self.db.record_metric_snapshot(777, views=1500, forwards=0, reactions=0)

        self._sync()

        ws = self._wb()["Продажи_рекламы"]
        self.assertEqual(ws["J6"].value, "Снято")
        self.assertEqual(ws["G6"].value, 1500)
        self.assertEqual(AdStatus.REMOVED.value, "removed")

    # ---------- общее ----------

    def test_buys_sheet_never_touched_and_nothing_written_when_no_changes(self):
        self._stat(1100, utc(2026, 9, 22, 8))
        book_ad(self.db, advertiser="A", price=10, scheduled_at=utc(2026, 9, 25), duration_hours=48)
        buys_before = [[c.value for c in row] for row in self._wb()["Закупы"].iter_rows()]

        self._sync()
        self.assertEqual([[c.value for c in row] for row in self._wb()["Закупы"].iter_rows()], buys_before)

        mtime = self.xlsx.stat().st_mtime_ns
        self.assertFalse(self._sync().changed)
        self.assertEqual(self.xlsx.stat().st_mtime_ns, mtime)  # без изменений файл не переписывается

    def test_empty_db_changes_nothing(self):
        result = self._sync()
        self.assertFalse(result.changed)

    def test_missing_file_raises_readable_error(self):
        with self.assertRaises(TrackerSyncError):
            sync_tracker(self.db, self.xlsx.with_name("nope.xlsx"))


if __name__ == "__main__":
    unittest.main()
