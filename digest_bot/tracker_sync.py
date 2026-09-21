"""Синхронизация SQLite -> tracker/channel_tracker.xlsx.

Пишет только значения в строки-заготовки, которые уже есть в листах «Метрики»
и «Продажи_рекламы» (формулы в них не трогаются, новые строки не создаются).
Лист «Закупы» не затрагивается вообще — его ведёт человек.

Синхронизация идемпотентна: день (по дате) или бронь (по id) второй раз в
таблицу не попадают. Если файл занят (открыт в Excel) — сохранение не
произойдёт, а данные останутся в БД и уедут в таблицу при следующем запуске.
"""
from __future__ import annotations

import logging
import os
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, time, timezone, tzinfo
from pathlib import Path
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from .db import Database
from .models import AdBooking, AdStatus

logger = logging.getLogger(__name__)

METRICS_SHEET = "Метрики"
SALES_SHEET = "Продажи_рекламы"

FIRST_DATA_ROW = 5
HEADER_ROW = 4

# Метрики: A Дата · B Подписчики · C Прирост (формула) · D Постов · E Охват · F ERR (формула)
M_DATE, M_SUBS, M_POSTS, M_REACH, M_FORMULA_COL = 1, 2, 4, 5, 6

# Продажи_рекламы: A Дата · B Рекламодатель · C Контакт · D Формат · E Цена · F Валюта ·
# G Просмотры · H CPM (формула) · I ERID · J Статус · K Комментарий · L ID брони (служебная)
S_DATE, S_ADVERTISER, S_CONTACT, S_FORMAT, S_PRICE, S_CURRENCY = 1, 2, 3, 4, 5, 6
S_VIEWS, S_FORMULA_COL, S_ERID, S_STATUS, S_NOTES, S_ID = 7, 8, 9, 10, 11, 12
S_ID_HEADER = "ID брони"

AD_STATUS_RU = {
    AdStatus.BOOKED: "Забронировано",
    AdStatus.PUBLISHED: "Опубликовано",
    AdStatus.REMOVED: "Снято",
    AdStatus.PAID: "Оплачено",
    AdStatus.CANCELLED: "Отменено",
}

_EXAMPLE_FILL_RGB = "EAF3FF"  # как EXAMPLE_FILL в tracker/build_tracker.py

# синий текст = «вписано человеком/ботом», как в легенде трекера
_INPUT_FONT = Font(name="Arial", color="0000FF", size=11)


class TrackerSyncError(RuntimeError):
    pass


@dataclass
class SyncResult:
    metrics_added: int = 0
    ads_added: int = 0
    skipped_no_room: int = 0  # записи, для которых кончились строки-заготовки

    @property
    def changed(self) -> bool:
        return bool(self.metrics_added or self.ads_added)


def _local_tz(name: str) -> tzinfo:
    try:
        return ZoneInfo(name)
    except (ZoneInfoNotFoundError, ValueError):
        logger.warning("Часовой пояс %r не найден (на Windows нужен пакет tzdata) — использую UTC", name)
        return timezone.utc


def _to_local_date(dt: datetime, tz: tzinfo) -> date:
    if dt.tzinfo is None:  # в БД naive-время — это UTC (datetime.utcnow())
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(tz).date()


def _cell_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _is_example_row(ws: Worksheet, row: int) -> bool:
    """Строка-пример шаблона (голубая заливка) — не настоящие данные."""
    fill = ws.cell(row=row, column=1).fill
    rgb = fill.fgColor.rgb if fill and fill.fill_type == "solid" else None
    return isinstance(rgb, str) and rgb.upper().endswith(_EXAMPLE_FILL_RGB)


def _is_formula(value) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _next_empty_row(ws: Worksheet, formula_col: int, key_cols: tuple[int, ...]) -> int | None:
    """Первая строка-заготовка (в ней есть формула) с пустыми ключевыми колонками."""
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        if not _is_formula(ws.cell(row=r, column=formula_col).value):
            continue
        if all(_is_blank(ws.cell(row=r, column=c).value) for c in key_cols):
            return r
    return None


def _put(ws: Worksheet, row: int, col: int, value) -> None:
    cell = ws.cell(row=row, column=col)
    cell.value = value
    cell.font = copy(_INPUT_FONT)


def _require_sheet(wb, name: str) -> Worksheet:
    if name not in wb.sheetnames:
        raise TrackerSyncError(f"В файле трекера нет листа «{name}»")
    return wb[name]


# ---------- Метрики ----------


def _daily_metrics(db: Database, tz: tzinfo) -> dict[date, dict]:
    """{локальная дата: {subscribers, posts, reach}} по всем дням, где есть замер подписчиков."""
    days: dict[date, dict] = {}
    for stat in db.all_channel_stats():  # по возрастанию времени — последний замер дня перезапишет предыдущие
        day = _to_local_date(datetime.fromisoformat(stat["captured_at"]), tz)
        days[day] = {"subscribers": stat["subscribers"], "posts": 0, "reach": None}

    views_by_day: dict[date, list[int]] = {}
    for draft in db.all_published():
        if draft.published_at is None:
            continue
        day = _to_local_date(draft.published_at, tz)
        if day not in days:
            continue
        days[day]["posts"] += 1
        if draft.channel_message_id:
            metric = db.latest_metric(draft.channel_message_id)
            if metric and metric["views"] is not None:
                views_by_day.setdefault(day, []).append(metric["views"])

    for day, views in views_by_day.items():
        days[day]["reach"] = round(sum(views) / len(views))
    return days


def _sync_metrics(ws: Worksheet, db: Database, tz: tzinfo, result: SyncResult) -> None:
    # строки-примеры не считаем: у шаблона пример датирован «сегодня» и заблокировал бы этот день
    present = {
        d for r in range(FIRST_DATA_ROW, ws.max_row + 1)
        if not _is_example_row(ws, r)
        and (d := _cell_date(ws.cell(row=r, column=M_DATE).value)) is not None
    }
    days = _daily_metrics(db, tz)
    for day in sorted(d for d in days if d not in present):
        row = _next_empty_row(ws, M_FORMULA_COL, (M_DATE, M_SUBS))
        if row is None:
            result.skipped_no_room += 1
            continue
        info = days[day]
        _put(ws, row, M_DATE, datetime.combine(day, time()))
        _put(ws, row, M_SUBS, info["subscribers"])
        _put(ws, row, M_POSTS, info["posts"])
        if info["reach"] is not None:
            _put(ws, row, M_REACH, info["reach"])
        result.metrics_added += 1


# ---------- Продажи_рекламы ----------


def _existing_ad_ids(ws: Worksheet) -> set[int]:
    ids = set()
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        value = ws.cell(row=r, column=S_ID).value
        try:
            ids.add(int(value))
        except (TypeError, ValueError):
            continue
    return ids


def _ensure_id_header(ws: Worksheet) -> None:
    header = ws.cell(row=HEADER_ROW, column=S_ID)
    if not _is_blank(header.value):
        return
    header.value = S_ID_HEADER
    style_src = ws.cell(row=HEADER_ROW, column=S_NOTES)
    header._style = copy(style_src._style)
    ws.column_dimensions["L"].width = 11


def _ad_views(db: Database, ad: AdBooking) -> int | None:
    if not ad.channel_message_id:
        return None
    metric = db.latest_metric(ad.channel_message_id)
    return metric["views"] if metric else None


def _sync_ads(ws: Worksheet, db: Database, tz: tzinfo, result: SyncResult) -> None:
    known = _existing_ad_ids(ws)
    for ad in db.all_ads():
        if ad.id in known:
            continue
        row = _next_empty_row(ws, S_FORMULA_COL, (S_DATE, S_ADVERTISER, S_PRICE))
        if row is None:
            result.skipped_no_room += 1
            continue
        _ensure_id_header(ws)
        _put(ws, row, S_DATE, datetime.combine(_to_local_date(ad.scheduled_at, tz), time()))
        _put(ws, row, S_ADVERTISER, ad.advertiser)
        _put(ws, row, S_CONTACT, ad.contact or None)
        _put(ws, row, S_FORMAT, f"Пост {ad.duration_hours}ч")
        _put(ws, row, S_PRICE, ad.price)
        _put(ws, row, S_CURRENCY, ad.currency)
        views = _ad_views(db, ad)
        if views is not None:
            _put(ws, row, S_VIEWS, views)
        _put(ws, row, S_ERID, ad.erid or None)
        _put(ws, row, S_STATUS, AD_STATUS_RU[ad.status])
        _put(ws, row, S_NOTES, ad.notes or None)
        _put(ws, row, S_ID, ad.id)
        result.ads_added += 1


# ---------- точка входа ----------


def sync_tracker(db: Database, xlsx_path: str | Path, timezone_name: str = "UTC") -> SyncResult:
    path = Path(xlsx_path)
    if not path.exists():
        raise TrackerSyncError(f"Файл трекера не найден: {path}")

    tz = _local_tz(timezone_name)
    wb = load_workbook(path)
    metrics_ws = _require_sheet(wb, METRICS_SHEET)
    sales_ws = _require_sheet(wb, SALES_SHEET)

    result = SyncResult()
    _sync_metrics(metrics_ws, db, tz, result)
    _sync_ads(sales_ws, db, tz, result)

    if result.skipped_no_room:
        logger.warning(
            "В трекере закончились строки-заготовки: %s записей не внесено. "
            "Расширьте формулы вниз или перегенерируйте лист.", result.skipped_no_room,
        )
    if not result.changed:
        return result

    # формулы без кэшированных значений — Excel пересчитает их при открытии
    wb.calculation.fullCalcOnLoad = True
    # сначала во временный файл: обрыв посреди записи не должен портить трекер
    tmp = path.with_name(path.stem + ".tmp.xlsx")
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError as exc:
        tmp.unlink(missing_ok=True)
        raise TrackerSyncError(
            f"Не удалось сохранить {path} — файл открыт в Excel? Закройте его; данные не потеряны, "
            "они попадут в таблицу при следующем запуске."
        ) from exc
    except Exception:
        tmp.unlink(missing_ok=True)
        raise
    return result
